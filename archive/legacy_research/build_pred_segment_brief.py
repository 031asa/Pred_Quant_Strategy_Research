from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
OUTPUT = BASE_DIR / "14点30分段_Pred修正方案分析.docx"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "6B7280"
GREEN = "E8F3EA"
AMBER = "FFF4D6"
RED = "9B1C1C"
BLACK = "111827"


def set_font(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, data, widths_dxa, percent_cols=None, highlight_rows=None):
    percent_cols = percent_cols or set()
    highlight_rows = highlight_rows or {}
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    header = table.rows[0]
    repeat_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(text)), size=9.2, bold=True, color="FFFFFF")

    for row_index, values in enumerate(data):
        row = table.add_row()
        fill = highlight_rows.get(row_index, "FFFFFF" if row_index % 2 == 0 else LIGHT_GRAY)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_shading(cell, fill)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            text = f"{value:.2%}" if col_index in percent_cols and isinstance(value, (int, float)) else str(value)
            color = RED if isinstance(value, str) and value.startswith("−") else BLACK
            set_font(p.add_run(text), size=9.2, color=color)
    set_table_geometry(table, widths_dxa)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=16 if level == 1 else 13, bold=True, color=BLUE if level <= 2 else NAVY)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), bold=True)
        set_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_font(p.add_run(text))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    set_font(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    set_font(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(f"{label}  "), size=11, bold=True, color=NAVY)
    set_font(p.add_run(text), size=11, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for style_name, size, before, after in (
    ("Heading 1", 16, 12, 6),
    ("Heading 2", 13, 10, 5),
    ("Heading 3", 12, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLUE if style_name != "Heading 3" else NAVY)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_after = Pt(0)
set_font(hp.add_run("模型验证简报  |  Pred 分时修正"), size=8.5, color=MID_GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_after = Pt(0)
set_font(fp.add_run("内部研究材料 · 2026-08-05"), size=8.5, color=MID_GRAY)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(10)
title.paragraph_format.space_after = Pt(5)
set_font(title.add_run("14:30 分段下的 Pred 修正方案分析"), size=23, bold=True, color=NAVY)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(12)
set_font(subtitle.add_run("详细说明版：从“数量对称”转向“空头精准率优先”"), size=13, color=MID_GRAY)

meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(14)
set_font(meta.add_run("数据：pred_eval.parquet  |  有效评估样本：1,670,400  |  滚动预热：20 个交易日"), size=9.5, color=MID_GRAY)

add_callout(
    doc,
    "当前判断",
    "14:30 前本身偏空且空头胜率高于随机水平，不应为了分布对称而上调 Pred。下一阶段应收紧做空选择、允许不交易区间；14:30 后可保留小幅单向去多头偏置，但不能追求多空各半。",
)

add_heading(doc, "阅读前先分清四个概念", level=2)
add_body(doc, "Pred / score：模型给每个样本的相对方向分数。分数较高更偏多，分数较低更偏空。", bold_prefix="Pred / score：")
add_body(doc, "分类阈值：把连续分数转换成多空方向的分界线。当前基准阈值为 0.5。", bold_prefix="分类阈值：")
add_body(doc, "做空占比：被判定为空单的样本数量占比。它只表示做空多少，不表示做空是否准确。", bold_prefix="做空占比：")
add_body(doc, "做空胜率：所有被判定为空单的样本中，Tag<0 的比例。它才是本阶段真正希望提高的指标。", bold_prefix="做空胜率：")
add_callout(doc, "关键区别", "空单变多不等于空单变准。修改 Pred 或阈值可以改变多空数量，但只有新增空单本身更容易下跌，整体做空胜率才会上升。", fill=AMBER)

add_heading(doc, "一、数据事实：14:30 确实形成两种不同状态")
add_body(doc, "在剔除最初 20 个交易日预热期后，14:30 前共有 1,461,600 个样本，14:30 及以后共有 208,800 个样本。两段的 score 分布中心、多空方向和胜率均有明显差异。")

baseline = [
    ["14:30 前", "1,461,600", "0.48309", 0.3635, 0.5017, 0.6365, 0.5111],
    ["14:30 及以后", "208,800", "0.54009", 0.6912, 0.5570, 0.3088, 0.5086],
]
add_table(
    doc,
    ["时段", "样本量", "滚动中心", "做多占比", "做多胜率", "做空占比", "做空胜率"],
    baseline,
    [1420, 1320, 1240, 1320, 1320, 1320, 1420],
    percent_cols={3, 4, 5, 6},
    highlight_rows={0: GREEN, 1: LIGHT_BLUE},
)

add_bullet(doc, "14:30 前：滚动中心约为 0.48309，做空占比 63.65%，做空胜率 51.11%，说明这一时段天然偏空。")
add_bullet(doc, "14:30 后：滚动中心约为 0.54009，做多占比 69.12%，做多胜率 55.70%，说明这一时段明显偏多。")
add_bullet(doc, "因此，一个覆盖全天的对称修正公式会把两种不同状态混在一起：它可能压低尾盘多头偏置，却同时破坏盘中已有的空头结构。")

add_heading(doc, "二、为什么整体下调 Pred 不等于提高空头胜率")
add_heading(doc, "2.1 下调 Pred，本质上是在移动分类阈值", level=2)
add_body(doc, "假设所有分数统一下调 Δ，再与 0.5 比较。数学上，这与保持原分数不变、把原始分界线从 0.5 提高到 0.5+Δ 完全等价。")
add_callout(doc, "等价关系", "adjusted_score = score − Δ；adjusted_score < 0.5 等价于 score < 0.5+Δ。也就是说，整体平移不会产生新的预测信息，只会改变哪些样本被划为空单。")
add_body(doc, "统一平移也不会改变样本排序。例如原分数为 0.40、0.45、0.52，全部减去 0.02 后变为 0.38、0.43、0.50，三者强弱顺序没有变化。因此 AUC 或排序能力不会因为整体平移而提高。")

add_heading(doc, "2.2 用 100 笔样本理解“空单增加但胜率下降”", level=2)
add_body(doc, "按照14:30前的真实比例，每100笔样本大约有64笔空单、36笔多单；现有空单胜率约51.11%，即64笔空单中大约33笔方向正确。")
add_body(doc, "如果下调 Pred，把5笔接近0.5的边界多单改为空单，原来的64笔空单并不会因此变得更准确，只是空单池扩大到69笔。只有这5笔新增空单的胜率高于原来的51.11%，整体空头胜率才会提高。")
add_callout(doc, "示意例子（不是实测值）", "如果新增5笔空单只有约45%的胜率，合并后的空头胜率会从约51.11%降到约50.7%。这就是“方向数量增加、方向质量下降”。", fill=AMBER)
add_body(doc, "因此，14:30前正确的问题不是“怎样再制造更多空单”，而是“Pred较低的样本是否真的比接近0.5的样本更容易下跌”。如果答案为是，就应该只保留更强的空头；如果答案为否，单靠平移 Pred 无法提高胜率。")

add_heading(doc, "三、已测试方案：单向修正能增加空单，但不能提高空单质量")
add_body(doc, "复核采用单向公式：只有过去 20 个交易日的滚动中位数高于 0.5 时才向下修正；低于 0.5 时保持原 Pred，不做反向上调。修正强度记为 ρ。")

effect = [
    ["0%", 0.5955, 0.5110, 0.3088, 0.5086, "—"],
    ["25%", 0.6014, 0.5109, 0.3557, 0.5087, "50.92%"],
    ["50%", 0.6076, 0.5109, 0.4048, 0.5082, "50.68%"],
    ["75%", 0.6139, 0.5106, 0.4547, 0.5062, "50.11%"],
    ["100%", 0.6203, 0.5103, 0.5049, 0.5030, "49.41%"],
]
add_table(
    doc,
    ["修正强度", "全体做空占比", "全体做空胜率", "尾盘做空占比", "尾盘做空胜率", "尾盘新增空单胜率"],
    effect,
    [1300, 1580, 1580, 1560, 1560, 1780],
    percent_cols={1, 2, 3, 4},
    highlight_rows={1: AMBER},
)

add_body(doc, "结果说明，方案在机械层面是有效的：修正越强，做空占比越高；但新增空单主要来自原来接近分类边界的样本，这些样本没有稳定的做空优势。")
add_bullet(doc, "14:30 前几乎不受影响：100% 修正后做空占比仍约 63.67%，做空胜率仍约 51.11%。这证明单向规则成功避免了错误上调。")
add_bullet(doc, "14:30 后从 25% 修正开始，做空占比由 30.88%升至 35.57%，整体做空胜率基本持平；继续增强后，做空胜率逐步降至 50.30%。")
add_bullet(doc, "被从多头改为空头的新增样本，其做空胜率由 25% 修正时的 50.92%降至完全修正时的 49.41%。追求多空各半会把低质量边界样本强行划为空头。")

add_callout(
    doc,
    "阶段结论",
    "25% 单向修正可以作为尾盘的保守候选，但其胜率改善只有约 0.01 个百分点，不构成明确优势。50% 以上修正应暂缓。当前更重要的问题不是让空单变多，而是筛出更可靠的空单。",
    fill=AMBER,
)

add_heading(doc, "四、改进方案：分时段、双阈值、允许不交易")
add_heading(doc, "4.1 14:30 前：空头精准率优先", level=2)
add_body(doc, "这一时段已经偏空，目标应从“增加空单数量”改为“保留更强的空头信号”。建议保持 score 的原始排序，并引入独立做空阈值 L：")
add_callout(doc, "方向规则", "score < L 时做空；score > U 时做多；L ≤ score ≤ U 时不交易。初始可固定 U=0.50，重点搜索 L<0.50。", fill=GREEN)
add_bullet(doc, "先测试 L=0.40～0.50，步长 0.01；同时用等样本分箱检查 score 越低时实际下跌率是否稳定提高。")
add_bullet(doc, "如果降低 L 后空头胜率上升、且不同月份均能复现，说明模型的空头排序有效，可以用覆盖率换精准率。")
add_bullet(doc, "如果低分样本的空头胜率没有提高，说明问题不在阈值，而在 Pred 对盘中空头缺少区分能力；统一平移无法解决。")

add_heading(doc, "4.2 14:30 后：小幅去多头偏置", level=2)
add_body(doc, "尾盘分布明显偏多，可继续保留“只向下、不向上”的单向修正。现有结果中 ρ=25% 对胜率伤害最小，可作为下一轮候选；但不应预设多空必须各占 50%。")
add_bullet(doc, "将 ρ 限定在 0%～30%的小范围内细化测试，例如每 5% 一个刻度。")
add_bullet(doc, "同时保留尾盘不交易区间，避免把 0.50 附近的模糊多头直接翻为空头。")
add_bullet(doc, "鉴于 Tag 在 14:30 后可能跨越收盘并连接下一交易日，应先确认标签周期，再解释尾盘效果。")

add_heading(doc, "五、验证方法与通过标准")
add_number(doc, "按 14:30 前后分别建表，不再使用一个全天阈值。")
add_number(doc, "阈值只能使用历史数据确定：用过去 20 或 60 个交易日估计，在下一交易日或下一月份验证。")
add_number(doc, "每个候选阈值同时报告空头胜率、空头样本量、覆盖率，以及按月份和合约拆分后的稳定性。")
add_number(doc, "先比较胜率和方向；只有样本外胜率稳定改善后，再计算收益率、交易成本和累计曲线。")

criteria = [
    ["主要指标", "14:30 前空头胜率", "相对 0.5 基准稳定提高，而不是只在全样本内提高"],
    ["覆盖约束", "空头样本量与覆盖率", "不能通过只保留极少样本制造虚高胜率"],
    ["稳定性", "按月、按合约", "改善不能集中在少数月份或少数合约"],
    ["尾盘约束", "新增空单胜率", "至少不低于原尾盘空单胜率，且不随修正强度持续下降"],
]
add_table(
    doc,
    ["检查维度", "观察指标", "建议通过标准"],
    criteria,
    [1520, 2480, 5360],
)

add_heading(doc, "六、建议的下一步")
add_body(doc, "优先生成“14:30 前做空阈值扫描表”，把 L 从 0.40 扫描至 0.50，并同步计算空头胜率、样本量和覆盖率；随后采用滚动历史窗口做样本外验证。14:30 后暂以 ρ=25% 作为对照组，不继续强化修正。")
add_body(doc, "这套方案的核心不是修饰 Pred 的外观，而是明确区分两个目标：分布偏置修正负责控制方向数量，阈值和不交易区间负责提高信号质量。只有后者在样本外成立，才说明策略真正得到改善。")

doc.core_properties.title = "14:30分段下的Pred修正方案分析"
doc.core_properties.subject = "Pred分时阈值与空头精准率验证"
doc.core_properties.author = "策略研究"
doc.core_properties.keywords = "Pred, 14:30, 阈值, 空头胜率, 分时分析"
doc.save(OUTPUT)
print(OUTPUT)
